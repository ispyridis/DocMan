import os
import re
from fpdf import FPDF
from docx import Document
from langdetect import detect
from collections import OrderedDict
# Update this path to where your templates are stored
template_directory = os.path.join(os.path.dirname(__file__), 'Templates')

def list_templates(template_directory):
    """
    List all .txt template files in the specified directory and its subdirectories.
    Includes relative path for templates in subdirectories to enable selecting templates from different contexts.
    """
    templates = []
    for root, dirs, files in os.walk(template_directory):
        for file in files:
            if file.endswith('.txt'):
                relative_path = os.path.relpath(os.path.join(root, file), template_directory)
                templates.append(relative_path)
    return templates


def parse_template(content):
    """
    Extract unique placeholders within brackets from the given content string.
    """
    placeholders = re.findall(r'\[(.*?)\]', content)
    unique_placeholders = list(OrderedDict.fromkeys(placeholders))

    return unique_placeholders, content

class PDF(FPDF):
    def __init__(self, font_mappings):
        super().__init__()
        self.font_mappings = font_mappings

    def add_fonts(self, content: str):
        lang = detect(content)
        font_name, font_path = self.font_mappings.get(lang, self.font_mappings['en'])
        if not os.path.exists(font_path):
            raise Exception(f"Font file {font_path} does not exist!")
        self.add_font(font_name, '', font_path, uni=True)
        self.set_font(font_name, '', 12)

def save_filled_template(content, filename, format='txt', export_dir = os.path.join(os.path.dirname(__file__), 'exported_docs')):
    font_mappings = {
        'en': ('DejaVu', os.path.join('fonts', 'DejaVuSansCondensed.ttf')),
        'el': ('DejaVu', os.path.join('fonts', 'DejaVuSansCondensed.ttf')),
    }

    os.makedirs(export_dir, exist_ok=True)
    filepath = os.path.join(export_dir, f"{filename}.{format}")
    
    if format == 'pdf':
        pdf = PDF(font_mappings)
        pdf.add_page()
        pdf.add_fonts(content)
        pdf.multi_cell(0, 10, content)
        pdf.output(filepath)
    elif format == 'docx':
        doc = Document()
        # Insert more complex formatting here if needed
        doc.add_paragraph(content)
        doc.save(filepath)
    else:
        with open(filepath, 'w', encoding='utf-8') as file:
            file.write(content)
